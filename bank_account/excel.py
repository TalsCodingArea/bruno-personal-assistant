"""Read and normalize the first sheet of an Israeli bank Excel export."""

import hashlib
import json
import re
import unicodedata
from collections.abc import Mapping, Sequence
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from bank_account.models import BankTransaction, TransactionDirection

_MONEY_PLACES = Decimal("0.01")
_DIRECTIONAL_MARKS = dict.fromkeys(map(ord, "\u200e\u200f\u202a\u202b\u202c\u202d\u202e"), None)

_HEADER_ALIASES: Mapping[str, frozenset[str]] = {
    "date": frozenset(("תאריך", "תאריך עסקה")),
    "action": frozenset(("הפעולה", "פעולה", "סוג פעולה")),
    "details": frozenset(("פרטים", "פירוט", "תיאור")),
    "reference": frozenset(("אסמכתא", "אסמכתה", "מספר אסמכתא", "מס אסמכתא")),
    "debit": frozenset(("חובה", "חיוב")),
    "credit": frozenset(("זכות", "זיכוי")),
    "balance": frozenset(("יתרה בשח", "יתרה", "יתרה לאחר פעולה")),
    "value_date": frozenset(("תאריך ערך",)),
    "beneficiary": frozenset(("לטובת", "מוטב")),
    "purpose": frozenset(("עבור", "מטרה")),
}
_REQUIRED_HEADERS = frozenset(("date", "action", "debit", "credit", "balance"))
_DATE_FORMATS = ("%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%Y-%m-%d")


class BankExportError(ValueError):
    """Raised when an export cannot be mapped without risking bad finance data."""


def read_bank_export(path: Path, *, header_scan_rows: int = 30) -> tuple[BankTransaction, ...]:
    """Parse transactions from the first worksheet of an ``.xlsx``/``.xlsm`` file."""

    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise BankExportError(f"Unsupported Excel type {path.suffix!r}; use .xlsx or .xlsm")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        header_row, columns = _find_headers(worksheet.iter_rows(values_only=True), header_scan_rows)
        return _read_rows(worksheet, header_row, columns, workbook.epoch)
    finally:
        workbook.close()


def _find_headers(
    rows: Sequence[Sequence[object]] | Any, scan_limit: int
) -> tuple[int, dict[str, int]]:
    best_columns: dict[str, int] = {}
    for row_number, row in enumerate(rows, start=1):
        if row_number > scan_limit:
            break
        columns = _map_headers(row)
        if len(columns) > len(best_columns):
            best_columns = columns
        if _REQUIRED_HEADERS.issubset(columns):
            return row_number, columns
    found = ", ".join(sorted(best_columns)) or "none"
    missing = ", ".join(sorted(_REQUIRED_HEADERS - best_columns.keys()))
    raise BankExportError(
        f"Could not identify the bank header row in the first {scan_limit} rows; "
        f"found canonical columns: {found}; missing: {missing}"
    )


def _map_headers(row: Sequence[object]) -> dict[str, int]:
    aliases = {
        _normalize_header(alias): canonical
        for canonical, values in _HEADER_ALIASES.items()
        for alias in values
    }
    result: dict[str, int] = {}
    for index, value in enumerate(row):
        canonical = aliases.get(_normalize_header(value))
        if canonical is not None and canonical not in result:
            result[canonical] = index
    return result


def _read_rows(
    worksheet: Any, header_row: int, columns: Mapping[str, int], epoch: datetime
) -> tuple[BankTransaction, ...]:
    transactions: list[BankTransaction] = []
    errors: list[str] = []
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        values = {name: _cell(row, index) for name, index in columns.items()}
        if not _looks_like_transaction(values):
            continue
        try:
            transactions.append(_transaction(values, row_number=row_number, epoch=epoch))
        except BankExportError as exc:
            errors.append(str(exc))
    if errors:
        preview = "; ".join(errors[:5])
        suffix = f"; and {len(errors) - 5} more" if len(errors) > 5 else ""
        raise BankExportError(f"Invalid transaction rows: {preview}{suffix}")
    if not transactions:
        raise BankExportError("The first worksheet contains no transaction rows")
    return tuple(transactions)


def _transaction(
    values: Mapping[str, object], *, row_number: int, epoch: datetime
) -> BankTransaction:
    occurred_on = _parse_date(values.get("date"), epoch=epoch, field="date", row=row_number)
    debit = _parse_money(values.get("debit"), field="debit", row=row_number, optional=True)
    credit = _parse_money(values.get("credit"), field="credit", row=row_number, optional=True)
    debit_value = abs(debit or Decimal(0))
    credit_value = abs(credit or Decimal(0))
    if debit_value and credit_value:
        raise BankExportError(f"row {row_number}: both debit and credit contain amounts")
    if not debit_value and not credit_value:
        raise BankExportError(f"row {row_number}: neither debit nor credit contains an amount")
    direction = TransactionDirection.NEGATIVE if debit_value else TransactionDirection.POSITIVE
    amount = debit_value or credit_value
    balance = _parse_money(values.get("balance"), field="balance", row=row_number)
    if balance is None:
        raise AssertionError("required money parser returned None")

    action = _clean_text(values.get("action"))
    details = _clean_text(values.get("details"))
    beneficiary = _clean_text(values.get("beneficiary"))
    purpose = _clean_text(values.get("purpose"))
    title = next(
        (value for value in (details, purpose, beneficiary, action) if value),
        "Transaction",
    )
    description = _description(values, epoch=epoch)
    fingerprint = _fingerprint(values, occurred_on, debit_value, credit_value, balance, epoch)
    reference = _reference_text(values.get("reference"))
    return BankTransaction(
        fingerprint=fingerprint,
        uid=f"{reference}:{fingerprint}" if reference else f"hash:{fingerprint}",
        title=title[:200],
        occurred_on=occurred_on,
        direction=direction,
        amount=amount.quantize(_MONEY_PLACES, rounding=ROUND_HALF_UP),
        balance_after=balance.quantize(_MONEY_PLACES, rounding=ROUND_HALF_UP),
        description=description[:2000],
        action=action[:2000],
    )


def _looks_like_transaction(values: Mapping[str, object]) -> bool:
    return any(_has_value(values.get(name)) for name in ("date", "debit", "credit", "balance"))


def _has_value(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _cell(row: Sequence[object], index: int) -> object:
    return row[index] if index < len(row) else None


def _normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", _clean_text(value)).translate(_DIRECTIONAL_MARKS)
    text = text.replace('"', "").replace("'", "").replace("\u05f3", "").replace("\u05f4", "")
    return re.sub(r"[^\w\u0590-\u05ff]+", " ", text, flags=re.UNICODE).strip().casefold()


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).translate(_DIRECTIONAL_MARKS).split())


def _reference_text(value: object) -> str:
    """Keep bank references stable when Excel loads an integer as a float."""

    if isinstance(value, bool):
        return _clean_text(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, Decimal) and value == value.to_integral_value():
        return str(value.to_integral_value())
    return _clean_text(value)


def _parse_date(value: object, *, epoch: datetime, field: str, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float) and not isinstance(value, bool):
        converted = from_excel(value, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    text = _clean_text(value)
    for format_ in _DATE_FORMATS:
        try:
            return datetime.strptime(text, format_).date()
        except ValueError:
            continue
    raise BankExportError(f"row {row}: {field} is not a recognized date: {text!r}")


def _parse_money(value: object, *, field: str, row: int, optional: bool = False) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if optional:
            return None
        raise BankExportError(f"row {row}: {field} is blank")
    if isinstance(value, bool):
        raise BankExportError(f"row {row}: {field} is not numeric")
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    text = _clean_text(value).replace("₪", "").replace(" ", "")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    if "," in text and "." in text:
        decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
        thousands_separator = "." if decimal_separator == "," else ","
        text = text.replace(thousands_separator, "").replace(decimal_separator, ".")
    elif "," in text:
        parts = text.split(",")
        text = "".join(parts) if len(parts[-1]) == 3 else ".".join(parts)
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise BankExportError(f"row {row}: {field} is not numeric: {value!r}") from exc
    return -result if negative else result


def _description(values: Mapping[str, object], *, epoch: datetime) -> str:
    parts: list[str] = []
    labels = (
        ("details", "פרטים"),
        ("reference", "אסמכתא"),
        ("value_date", "תאריך ערך"),
        ("beneficiary", "לטובת"),
        ("purpose", "עבור"),
    )
    for name, label in labels:
        value = values.get(name)
        if not _has_value(value):
            continue
        if name == "value_date":
            try:
                rendered = _parse_date(value, epoch=epoch, field=name, row=0).isoformat()
            except BankExportError:
                rendered = _clean_text(value)
        elif name == "reference":
            rendered = _reference_text(value)
        else:
            rendered = _clean_text(value)
        parts.append(f"{label}: {rendered}")
    return " | ".join(parts)


def _fingerprint(
    values: Mapping[str, object],
    occurred_on: date,
    debit: Decimal,
    credit: Decimal,
    balance: Decimal,
    epoch: datetime,
) -> str:
    value_date = values.get("value_date")
    try:
        normalized_value_date = (
            _parse_date(value_date, epoch=epoch, field="value_date", row=0).isoformat()
            if _has_value(value_date)
            else ""
        )
    except BankExportError:
        normalized_value_date = _clean_text(value_date)
    payload = {
        "date": occurred_on.isoformat(),
        "action": _clean_text(values.get("action")),
        "details": _clean_text(values.get("details")),
        "reference": _reference_text(values.get("reference")),
        "debit": str(debit),
        "credit": str(credit),
        "balance": str(balance),
        "value_date": normalized_value_date,
        "beneficiary": _clean_text(values.get("beneficiary")),
        "purpose": _clean_text(values.get("purpose")),
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()
