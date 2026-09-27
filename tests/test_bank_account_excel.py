"""Excel parsing tests for the standalone bank-account importer."""

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from bank_account.excel import BankExportError, read_bank_export
from bank_account.models import TransactionDirection

HEADERS = [
    "תאריך",
    "הפעולה",
    "פרטים",
    "אסמכתא",
    "חובה",
    "זכות",
    "יתרה בש''ח",
    "תאריך ערך",
    "לטובת",
    "עבור",
]


def _workbook(path: Path, rows: list[list[object]], *, header_row: int = 5) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for _ in range(header_row - 1):
        sheet.append(["bank export"])
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.create_sheet("ignored").append(["not", "bank", "data"])
    workbook.save(path)


def test_reads_first_sheet_and_maps_debit_and_credit(tmp_path: Path) -> None:
    path = tmp_path / "bank.xlsx"
    _workbook(
        path,
        [
            [
                date(2026, 9, 2),
                "כרטיס אשראי",
                "קניות",
                "abc-1",
                125.4,
                None,
                5000,
                date(2026, 9, 2),
                "",
                "",
            ],
            [
                "03/09/2026",
                "משכורת",
                "שכר",
                "abc-2",
                None,
                "10,000.50",
                15000.5,
                "03/09/2026",
                "חברה",
                "ספטמבר",
            ],
        ],
    )

    rows = read_bank_export(path)

    assert len(rows) == 2
    assert rows[0].direction is TransactionDirection.NEGATIVE
    assert rows[0].amount == Decimal("125.40")
    assert rows[0].balance_after == Decimal("5000.00")
    assert rows[0].title == "קניות"
    assert rows[0].uid == f"abc-1:{rows[0].fingerprint}"
    assert "אסמכתא: abc-1" in rows[0].description
    assert rows[1].direction is TransactionDirection.POSITIVE
    assert rows[1].amount == Decimal("10000.50")
    assert rows[1].uid == f"abc-2:{rows[1].fingerprint}"


def test_discovers_shifted_header_and_reordered_columns(tmp_path: Path) -> None:
    path = tmp_path / "shifted.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["metadata"])
    sheet.append(["more metadata"])
    headers = ["זכות", 'יתרה בש"ח', "תאריך", "חובה", "הפעולה"]
    sheet.append(headers)
    sheet.append([250, 1250, "2026-09-10", None, "החזר"])
    workbook.save(path)

    rows = read_bank_export(path)

    assert rows[0].occurred_on == date(2026, 9, 10)
    assert rows[0].amount == Decimal("250.00")
    assert rows[0].action == "החזר"


def test_rejects_ambiguous_amount_without_partial_result(tmp_path: Path) -> None:
    path = tmp_path / "ambiguous.xlsx"
    _workbook(path, [["10/09/2026", "פעולה", "", "1", 10, 20, 100, "", "", ""]])

    with pytest.raises(BankExportError, match="both debit and credit"):
        read_bank_export(path)


def test_uses_hash_uid_when_reference_is_blank(tmp_path: Path) -> None:
    path = tmp_path / "no-reference.xlsx"
    _workbook(path, [["10/09/2026", "פעולה", "פרטים", "", 10, None, 100, "", "", ""]])

    transaction = read_bank_export(path)[0]

    assert transaction.uid == f"hash:{transaction.fingerprint}"


def test_reused_bank_reference_produces_unique_transaction_uids(tmp_path: Path) -> None:
    path = tmp_path / "reused-reference.xlsx"
    _workbook(
        path,
        [
            ["10/09/2026", "פעולה", "ראשון", "abc-1", 10, None, 100, "", "", ""],
            ["11/09/2026", "פעולה", "שני", "abc-1", 20, None, 80, "", "", ""],
        ],
    )

    first, second = read_bank_export(path)

    assert first.uid.startswith("abc-1:")
    assert second.uid.startswith("abc-1:")
    assert first.uid != second.uid
